from collections import defaultdict
from io import BytesIO
from uuid import UUID

from fastapi import HTTPException, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.emploi_du_temps import CreneauHoraire, SeanceCours
from app.models.parametrage import Classe
from app.models.personnel import Personnel
from app.schemas.emploi_du_temps import (
    JOURS_LABELS,
    ConflitEdt,
    ConflitsListResponse,
    CreneauCreate,
    CreneauResponse,
    CreneauUpdate,
    GrilleEdtResponse,
    LigneGrille,
    SeanceCreate,
    SeanceResponse,
    SeanceUpdate,
)
from app.services import parametrage_service


def _seance_options():
    return (
        selectinload(SeanceCours.creneau),
        selectinload(SeanceCours.matiere),
        selectinload(SeanceCours.personnel),
        selectinload(SeanceCours.classe),
    )


async def _get_annee_id(db: AsyncSession, annee_id: UUID | None) -> UUID:
    if annee_id:
        return annee_id
    annee = await parametrage_service.get_annee_active(db)
    if annee is None:
        raise HTTPException(status_code=422, detail="Aucune année scolaire active")
    return annee.id


async def list_creneaux(db: AsyncSession, annee_scolaire_id: UUID | None = None) -> list[CreneauHoraire]:
    annee_id = await _get_annee_id(db, annee_scolaire_id)
    result = await db.execute(
        select(CreneauHoraire)
        .where(CreneauHoraire.annee_scolaire_id == annee_id)
        .order_by(CreneauHoraire.ordre)
    )
    return list(result.scalars().all())


async def create_creneau(db: AsyncSession, data: CreneauCreate) -> CreneauHoraire:
    annee_id = await _get_annee_id(db, data.annee_scolaire_id)
    if data.heure_fin <= data.heure_debut:
        raise HTTPException(status_code=422, detail="L'heure de fin doit être après l'heure de début")
    creneau = CreneauHoraire(annee_scolaire_id=annee_id, **data.model_dump(exclude={"annee_scolaire_id"}))
    db.add(creneau)
    await db.commit()
    await db.refresh(creneau)
    return creneau


async def update_creneau(db: AsyncSession, creneau_id: UUID, data: CreneauUpdate) -> CreneauHoraire:
    creneau = await db.get(CreneauHoraire, creneau_id)
    if creneau is None:
        raise HTTPException(status_code=404, detail="Créneau introuvable")
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(creneau, field, value)
    if creneau.heure_fin <= creneau.heure_debut:
        raise HTTPException(status_code=422, detail="L'heure de fin doit être après l'heure de début")
    await db.commit()
    await db.refresh(creneau)
    return creneau


async def delete_creneau(db: AsyncSession, creneau_id: UUID) -> None:
    creneau = await db.get(CreneauHoraire, creneau_id)
    if creneau is None:
        raise HTTPException(status_code=404, detail="Créneau introuvable")
    await db.delete(creneau)
    await db.commit()


def _detect_conflits(seances: list[SeanceCours]) -> list[ConflitEdt]:
    conflits: list[ConflitEdt] = []
    by_enseignant: dict[tuple, list[SeanceCours]] = defaultdict(list)
    by_salle: dict[tuple, list[SeanceCours]] = defaultdict(list)

    for s in seances:
        key_e = (s.personnel_id, s.creneau_id, s.jour_semaine)
        by_enseignant[key_e].append(s)
        salle = s.salle or (s.classe.salle if s.classe else None)
        if salle:
            key_s = (salle.lower(), s.creneau_id, s.jour_semaine)
            by_salle[key_s].append(s)

    for group in by_enseignant.values():
        if len(group) > 1:
            conflits.append(
                ConflitEdt(
                    type="enseignant",
                    message=f"Enseignant en double le {JOURS_LABELS[group[0].jour_semaine]}",
                    seance_ids=[s.id for s in group],
                )
            )

    for group in by_salle.values():
        if len(group) > 1:
            conflits.append(
                ConflitEdt(
                    type="salle",
                    message=f"Salle en double le {JOURS_LABELS[group[0].jour_semaine]}",
                    seance_ids=[s.id for s in group],
                )
            )

    return conflits


async def _load_seances(
    db: AsyncSession,
    annee_id: UUID,
    *,
    classe_id: UUID | None = None,
    personnel_id: UUID | None = None,
) -> list[SeanceCours]:
    query = (
        select(SeanceCours)
        .options(*_seance_options())
        .where(SeanceCours.annee_scolaire_id == annee_id)
    )
    if classe_id:
        query = query.where(SeanceCours.classe_id == classe_id)
    if personnel_id:
        query = query.where(SeanceCours.personnel_id == personnel_id)

    result = await db.execute(query)
    return list(result.scalars().all())


def _build_grille(
    titre: str,
    creneaux: list[CreneauHoraire],
    seances: list[SeanceCours],
    *,
    filter_classe: UUID | None = None,
    filter_personnel: UUID | None = None,
) -> GrilleEdtResponse:
    index: dict[tuple, SeanceCours] = {}
    for s in seances:
        if filter_classe and s.classe_id != filter_classe:
            continue
        if filter_personnel and s.personnel_id != filter_personnel:
            continue
        index[(s.creneau_id, s.jour_semaine)] = s

    all_seances_for_conflicts = seances if not filter_personnel else seances
    conflits = _detect_conflits(all_seances_for_conflicts)

    lignes: list[LigneGrille] = []
    for creneau in creneaux:
        cellules: list[SeanceResponse | None] = []
        for jour in range(5):
            seance = index.get((creneau.id, jour))
            cellules.append(SeanceResponse.model_validate(seance) if seance else None)
        lignes.append(
            LigneGrille(
                creneau=CreneauResponse.model_validate(creneau),
                cellules=cellules,
            )
        )

    return GrilleEdtResponse(titre=titre, jours=JOURS_LABELS, lignes=lignes, conflits=conflits)


async def get_grille_classe(db: AsyncSession, classe_id: UUID) -> GrilleEdtResponse:
    classe = await db.get(Classe, classe_id)
    if classe is None:
        raise HTTPException(status_code=404, detail="Classe introuvable")
    annee_id = classe.annee_scolaire_id
    creneaux = await list_creneaux(db, annee_id)
    seances = await _load_seances(db, annee_id, classe_id=classe_id)
    all_seances = await _load_seances(db, annee_id)
    grille = _build_grille(f"EDT — {classe.nom}", creneaux, all_seances, filter_classe=classe_id)
    grille.conflits = _detect_conflits(all_seances)
    return grille


async def get_grille_enseignant(db: AsyncSession, personnel_id: UUID) -> GrilleEdtResponse:
    personnel = await db.get(Personnel, personnel_id)
    if personnel is None:
        raise HTTPException(status_code=404, detail="Enseignant introuvable")
    annee_id = await _get_annee_id(db, None)
    creneaux = await list_creneaux(db, annee_id)
    seances = await _load_seances(db, annee_id, personnel_id=personnel_id)
    all_seances = await _load_seances(db, annee_id)
    return _build_grille(
        f"EDT — {personnel.prenoms} {personnel.nom}",
        creneaux,
        all_seances,
        filter_personnel=personnel_id,
    )


async def list_conflits(db: AsyncSession, annee_scolaire_id: UUID | None = None) -> ConflitsListResponse:
    annee_id = await _get_annee_id(db, annee_scolaire_id)
    seances = await _load_seances(db, annee_id)
    items = _detect_conflits(seances)
    return ConflitsListResponse(items=items, total=len(items))


async def _check_seance_conflicts(
    db: AsyncSession,
    data: SeanceCreate | SeanceUpdate,
    annee_id: UUID,
    seance_id: UUID | None = None,
    existing: SeanceCours | None = None,
) -> None:
    personnel_id = getattr(data, "personnel_id", None) or (existing.personnel_id if existing else None)
    creneau_id = getattr(data, "creneau_id", None) or (existing.creneau_id if existing else None)
    jour = getattr(data, "jour_semaine", None)
    if jour is None and existing:
        jour = existing.jour_semaine
    salle = getattr(data, "salle", None)
    if salle is None and existing:
        salle = existing.salle

    if personnel_id is None or creneau_id is None or jour is None:
        return

    result = await db.execute(
        select(SeanceCours)
        .options(selectinload(SeanceCours.classe))
        .where(
            SeanceCours.annee_scolaire_id == annee_id,
            SeanceCours.personnel_id == personnel_id,
            SeanceCours.creneau_id == creneau_id,
            SeanceCours.jour_semaine == jour,
        )
    )
    for other in result.scalars().all():
        if seance_id and other.id == seance_id:
            continue
        raise HTTPException(
            status_code=409,
            detail=f"Conflit enseignant le {JOURS_LABELS[jour]} à ce créneau",
        )

    if salle:
        result_salle = await db.execute(
            select(SeanceCours)
            .where(
                SeanceCours.annee_scolaire_id == annee_id,
                SeanceCours.creneau_id == creneau_id,
                SeanceCours.jour_semaine == jour,
            )
        )
        for other in result_salle.scalars().all():
            if seance_id and other.id == seance_id:
                continue
            other_salle = other.salle
            if other_salle and other_salle.lower() == salle.lower():
                raise HTTPException(
                    status_code=409,
                    detail=f"Conflit salle le {JOURS_LABELS[jour]} à ce créneau",
                )


async def create_seance(db: AsyncSession, data: SeanceCreate) -> SeanceCours:
    annee_id = await _get_annee_id(db, data.annee_scolaire_id)
    classe = await db.get(Classe, data.classe_id)
    if classe is None:
        raise HTTPException(status_code=404, detail="Classe introuvable")

    salle = data.salle or classe.salle
    await _check_seance_conflicts(db, data, annee_id)

    seance = SeanceCours(
        annee_scolaire_id=annee_id,
        salle=salle,
        **data.model_dump(exclude={"annee_scolaire_id", "salle"}),
    )
    db.add(seance)
    try:
        await db.commit()
    except Exception:
        await db.rollback()
        raise HTTPException(status_code=409, detail="Créneau déjà occupé pour cette classe") from None

    result = await db.execute(
        select(SeanceCours).options(*_seance_options()).where(SeanceCours.id == seance.id)
    )
    return result.scalar_one()


async def update_seance(db: AsyncSession, seance_id: UUID, data: SeanceUpdate) -> SeanceCours:
    result = await db.execute(
        select(SeanceCours).options(*_seance_options()).where(SeanceCours.id == seance_id)
    )
    seance = result.scalar_one_or_none()
    if seance is None:
        raise HTTPException(status_code=404, detail="Séance introuvable")

    merged = SeanceCreate(
        classe_id=seance.classe_id,
        creneau_id=data.creneau_id or seance.creneau_id,
        jour_semaine=data.jour_semaine if data.jour_semaine is not None else seance.jour_semaine,
        matiere_id=data.matiere_id or seance.matiere_id,
        personnel_id=data.personnel_id or seance.personnel_id,
        salle=data.salle if data.salle is not None else seance.salle,
    )
    await _check_seance_conflicts(db, merged, seance.annee_scolaire_id, seance_id=seance_id, existing=seance)

    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(seance, field, value)

    await db.commit()
    db.expire_all()
    result = await db.execute(
        select(SeanceCours)
        .options(*_seance_options())
        .where(SeanceCours.id == seance_id)
        .execution_options(populate_existing=True)
    )
    return result.scalar_one()


async def delete_seance(db: AsyncSession, seance_id: UUID) -> None:
    seance = await db.get(SeanceCours, seance_id)
    if seance is None:
        raise HTTPException(status_code=404, detail="Séance introuvable")
    await db.delete(seance)
    await db.commit()


def generate_edt_excel(grille: GrilleEdtResponse) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Emploi du temps"

    header_fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.cell(1, 1, grille.titre).font = Font(bold=True, size=14)
    headers = ["Créneau"] + grille.jours
    for col, h in enumerate(headers, 1):
        cell = ws.cell(2, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, ligne in enumerate(grille.lignes, 3):
        c = ligne.creneau
        label = f"{c.libelle}\n{c.heure_debut.strftime('%H:%M')}-{c.heure_fin.strftime('%H:%M')}"
        ws.cell(row_idx, 1, label)
        for jour_idx, seance in enumerate(ligne.cellules):
            if seance and seance.matiere:
                text = seance.matiere.libelle
                if seance.personnel:
                    text += f"\n({seance.personnel.prenoms} {seance.personnel.nom})"
                if seance.salle:
                    text += f"\n[{seance.salle}]"
                ws.cell(row_idx, jour_idx + 2, text)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
